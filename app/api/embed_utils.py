# app/api/embed_utils.py

from sentence_transformers import SentenceTransformer
import faiss
import numpy as np
import math
from pdfminer.high_level import extract_text as extract_pdf_text
from docx import Document as DocxDocument
import openpyxl

# 1) Инициализируем модель и FAISS-индекс
MODEL_NAME = "all-MiniLM-L6-v2"
model = SentenceTransformer(MODEL_NAME)
dim = model.get_sentence_embedding_dimension()
index = faiss.IndexFlatL2(dim)  # в памяти

# 2) Вспомогательная функция: извлечение текста
def extract_text_from_file(path: str, content: bytes) -> str:
    if path.lower().endswith(".pdf"):
        # сохраняем временно и читаем
        with open("temp.pdf", "wb") as f:
            f.write(content)
        return extract_pdf_text("temp.pdf")
    if path.lower().endswith(".docx"):
        doc = DocxDocument(path)
        return "\n".join(p.text for p in doc.paragraphs)
    if path.lower().endswith(".xlsx") or path.lower().endswith(".xlsm"):
        wb = openpyxl.load_workbook(path, data_only=True)
        text = []
        for sheet in wb.worksheets:
            for row in sheet.iter_rows(values_only=True):
                text.append(" ".join(str(c) for c in row if c))
        return "\n".join(text)
    # иначе пытаемся декодировать
    return content.decode("utf-8", errors="ignore")

# 3) Функция дробления на чанки ~500 токенов (приблизительно)
def chunk_text(text: str, max_tokens: int = 500) -> list[str]:
    words = text.split()
    chunks = []
    for i in range(0, len(words), max_tokens):
        chunk = " ".join(words[i : i + max_tokens])
        chunks.append(chunk)
    return chunks

# 4) Эмбеддим и индексируем массив чанков
def embed_and_index(chunks: list[str]) -> np.ndarray:
    vectors = model.encode(chunks, show_progress_bar=False, convert_to_numpy=True)
    index.add(vectors)
    return vectors

# 5) Поиск по запросу
def search_index(query: str, k: int = 5) -> tuple[np.ndarray, np.ndarray]:
    q_vec = model.encode([query], convert_to_numpy=True)
    distances, indices = index.search(q_vec, k)
    return distances[0], indices[0]
